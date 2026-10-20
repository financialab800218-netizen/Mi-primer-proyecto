import unicodedata
from almacenamiento import subir_archivo_a_nube
import os
from datetime import datetime
from docx import Document
from fpdf import FPDF
import PyPDF2

RUTA_DOCS = "/root/Mi-primer-proyecto/backend/app/static/docs"
URL_BASE = "https://lens-portside-masses.ngrok-free.dev/docs/"

def generar_contrato_tcp(nombre: str, nivel: str, categoria: str, porcentaje: float, bono_base: float):
    """Genera un contrato Word para un profesor TCP."""
    try:
        os.makedirs(RUTA_DOCS, exist_ok=True)
        
        doc = Document()
        doc.add_heading('CONTRATO DE LICENCIA DE CONTENIDO - TUTORIA CUBA', 0)
        
        doc.add_paragraph(f"En la ciudad de Bejucal, a los {datetime.now().day} días del mes de {datetime.now().month} de {datetime.now().year}.")
        
        doc.add_heading('PARTES', level=1)
        doc.add_paragraph("De una parte, la MIPYME 'Tutoria Cuba' representada por Yosbel Collazo Avila.")
        doc.add_paragraph(f"De otra parte, el Profesor TCP: {nombre}.")
        
        doc.add_heading('ACUERDOS', level=1)
        doc.add_paragraph(f"1. El profesor {nombre} prestará servicios de tutoría para el nivel {nivel}.")
        doc.add_paragraph(f"2. El profesor pertenece a la categoría {categoria}.")
        doc.add_paragraph(f"3. El porcentaje de reparto para el profesor será del {porcentaje}% sobre las ventas.")
        doc.add_paragraph(f"4. El bono base garantizado mensual es de {bono_base:,.2f} CUP.")
        
        doc.add_heading('CUMPLIMIENTO LEGAL', level=1)
        doc.add_paragraph("Este contrato cumple con el Decreto 160/2026 y la Ley 176. Tutoria Cuba no es una academia, es una plataforma tecnológica.")
        
        # Limpiar tildes y caracteres especiales del nombre
        nombre_limpio = unicodedata.normalize("NFKD", nombre).encode("ASCII", "ignore").decode("utf-8").replace(" ", "_")
        nombre_archivo = f"contrato_{nombre_limpio}.docx"
        ruta_completa = os.path.join(RUTA_DOCS, nombre_archivo)
        doc.save(ruta_completa)
        
        url_nube = subir_archivo_a_nube(nombre_archivo, nombre_archivo)
        return f"Contrato generado en la nube: {url_nube}"
    except Exception as e:
        return f"Error al generar contrato: {str(e)}"

def generar_pdf(texto: str, nombre_archivo: str):
    """Genera un archivo PDF a partir de un texto."""
    try:
        os.makedirs(RUTA_DOCS, exist_ok=True)
        
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        # Codificar el texto para soportar caracteres latinos
        texto_encoded = texto.encode('latin-1', 'replace').decode('latin-1')
        
        for linea in texto_encoded.split('\n'):
            pdf.cell(200, 10, txt=linea, ln=True)
            
        if not nombre_archivo.endswith('.pdf'):
            nombre_archivo += '.pdf'
            
        ruta_completa = os.path.join(RUTA_DOCS, nombre_archivo)
        pdf.output(ruta_completa)
        
        url_nube = subir_archivo_a_nube(nombre_archivo, nombre_archivo)
        return f"PDF generado en la nube: {url_nube}"
    except Exception as e:
        return f"Error al generar PDF: {str(e)}"

def leer_documento(nombre_archivo: str):
    """Lee un archivo Word o PDF y devuelve su texto."""
    ruta_completa = os.path.join(RUTA_DOCS, nombre_archivo)
    
    if not os.path.exists(ruta_completa):
        return f"Error: El archivo {nombre_archivo} no existe."

    try:
        if nombre_archivo.endswith('.docx'):
            doc = Document(ruta_completa)
            texto = "\n".join([parrafo.text for parrafo in doc.paragraphs])
            return f"Contenido del documento:\n\n{texto}"
            
        elif nombre_archivo.endswith('.pdf'):
            with open(ruta_completa, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                texto = ""
                for pagina in reader.pages:
                    texto += pagina.extract_text() + "\n"
            return f"Contenido del PDF:\n\n{texto}"
            
        else:
            return "Error: Formato no soportado. Use .docx o .pdf."
    except Exception as e:
        return f"Error al leer documento: {str(e)}"
